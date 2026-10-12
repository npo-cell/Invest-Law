#!/usr/bin/env python3
"""Upload des factsheets INVESTLAW vers invest-law.eu (section « Rapports de Gestion »).

Identifiants lus dans l'environnement : INVESTLAW_EMAIL, INVESTLAW_PASSWORD.

  login            envoie email + mot de passe ; le site envoie un code à 6 chiffres par email
  verify <code>    valide le code et enregistre la session (cookies)
  plan             associe chaque ligne du manifeste à une fiche (ISIN) et repère les rapports déjà présents
  upload [--include-flagged] [--limit N]
                   envoie les PDF du plan marqués « A_UPLOADER » ; n'en supprime aucun
"""
import base64, http.cookiejar, json, os, re, sys, unicodedata, urllib.error, urllib.request
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR = os.path.join(ROOT, 'INVESTLAW')
MANIFEST = os.path.join(DIR, 'MANIFEST_UPLOAD.xlsx')
STATE = os.path.join(os.environ.get('INVESTLAW_STATE_DIR', os.path.expanduser('~/.investlaw')), '')
BASE = 'https://invest-law.eu'

os.makedirs(STATE, exist_ok=True)
jar = http.cookiejar.LWPCookieJar(STATE + 'cookies.txt')
if os.path.exists(jar.filename):
    jar.load(ignore_discard=True, ignore_expires=True)
opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))


def api(method, path, body=None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(BASE + path, data=data, method=method,
                                 headers={'Content-Type': 'application/json', 'User-Agent': 'Mozilla/5.0'})
    try:
        with opener.open(req, timeout=120) as r:
            out = r.read()
            status = r.status
    except urllib.error.HTTPError as e:
        out, status = e.read(), e.code
    jar.save(ignore_discard=True, ignore_expires=True)
    try:
        return status, json.loads(out or b'{}')
    except ValueError:
        return status, {'raw': out[:300].decode(errors='replace')}


MOIS = ['janvier', 'fevrier', 'mars', 'avril', 'mai', 'juin', 'juillet', 'aout',
        'septembre', 'octobre', 'novembre', 'decembre']


def norm(s):
    return unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().lower().strip()


def name_ym(name):
    """'Rapport de gestion Aout 2026' -> '2026-08' ; None si le nom ne contient pas de mois et d'année."""
    m = re.search(r'(%s)\s+(\d{4})' % '|'.join(MOIS), norm(name))
    return '%s-%02d' % (m.group(2), MOIS.index(m.group(1)) + 1) if m else None


def rows():
    wb = openpyxl.load_workbook(MANIFEST)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    return wb, ws, hdr


def cmd_login():
    st, d = api('POST', '/api/login', {'username': os.environ['INVESTLAW_EMAIL'],
                                       'password': os.environ['INVESTLAW_PASSWORD']})
    if st == 200 and d.get('step') == 'verify':
        json.dump({'user_id': d.get('user_id'), 'verification_id': d.get('verification_id')},
                  open(STATE + 'verify.json', 'w'))
        print('CODE_ENVOYE', d.get('email', ''))
    else:
        print('ECHEC', st, d.get('message') or d)


def cmd_verify(code):
    ctx = json.load(open(STATE + 'verify.json'))
    st, d = api('POST', '/api/login/verify', dict(ctx, code=code))
    print('OK role=%s' % d.get('role') if st == 200 and d.get('success') else 'ECHEC %s %s' % (st, d.get('message') or d))


def cmd_plan():
    st, funds = api('GET', '/api/funds')
    if st != 200:
        sys.exit('Non connecté (%s) : relancer login/verify' % st)
    by_isin = {}
    for f in funds:
        if f.get('ISIN'):
            by_isin.setdefault(f['ISIN'].strip().upper(), []).append(f)
    wb, ws, hdr = rows()
    for col in ('FundId', 'Fiche (nom sur le site)', 'Plan', 'Détail'):
        if col not in hdr:
            ws.cell(1, len(hdr) + 1, col); hdr.append(col)
    H = {h: i for i, h in enumerate(hdr)}
    for r in ws.iter_rows(min_row=2):
        isin = r[H['ISIN']].value
        if r[H['Statut upload']].value == 'UPLOADE':
            continue
        cands = by_isin.get(isin, [])
        if len(cands) != 1:
            plan, det, fid, nom = ('FONDS_ABSENT', '', '', '') if not cands else ('ISIN_EN_DOUBLE', ','.join(str(c['FundId']) for c in cands), '', '')
        else:
            f = cands[0]; fid, nom = f['FundId'], f.get('Name', '')
            st, reps = api('GET', '/api/funds/%s/reports' % fid)
            reps = reps if isinstance(reps, list) else []
            d = r[H['Date (jj/mm/aaaa)']].value  # jj/mm/aaaa
            ym = d[6:10] + '-' + d[3:5]
            # report_date sur le site = date d'upload : on compare le mois/année lu dans le nom du rapport
            # tout document sauf les DIC/KID (« Rapport de gestion », « Fiche mensuelle », « Fiche produit »…)
            raps = [x for x in reps if not re.search(r'\b(dic|kid)\b|informations cles', norm(x.get('report_name')))]
            months = [(x, name_ym(x.get('report_name'))) for x in raps]
            same = [x for x, m in months if m == ym or norm(x.get('report_name')) == norm(r[H['Nom du rapport']].value)]
            newer = [x for x, m in months if m and m > ym]
            # nom sans mois : déposé au plus tard le mois de notre factsheet => forcément pas plus récent
            vague = [x for x, m in months if not m and str(x.get('report_date'))[:7] > ym]
            plan = ('DEJA_PRESENT' if same else 'PLUS_RECENT_PRESENT' if newer
                    else 'A_VERIFIER' if vague else 'A_UPLOADER')
            det = '; '.join('%s (%s)' % (x.get('report_name'), str(x.get('report_date'))[:10])
                            for x in same or newer or vague) or '%d rapport(s) existant(s)' % len(reps)
        r[H['FundId']].value, r[H['Fiche (nom sur le site)']].value, r[H['Plan']].value, r[H['Détail']].value = fid, nom, plan, det
    wb.save(MANIFEST)
    from collections import Counter
    print(Counter(r[H['Plan']].value for r in ws.iter_rows(min_row=2)))


def cmd_upload(include_flagged=False, limit=None):
    wb, ws, hdr = rows()
    H = {h: i for i, h in enumerate(hdr)}
    n = 0
    for r in ws.iter_rows(min_row=2):
        if r[H['Plan']].value != 'A_UPLOADER' or r[H['Statut upload']].value == 'UPLOADE':
            continue
        if r[H['A contrôler']].value and not include_flagged:
            continue
        if limit is not None and n >= limit:
            break
        path = os.path.join(DIR, r[H['Fichier (chemin dans INVESTLAW)']].value)
        isin = r[H['ISIN']].value
        if not os.path.basename(path).startswith(isin + '__'):
            r[H['Statut upload']].value = 'ERREUR: fichier/ISIN incohérent'; continue
        d = r[H['Date (jj/mm/aaaa)']].value
        st, resp = api('POST', '/api/funds/%s/reports' % r[H['FundId']].value, {
            'report_name': r[H['Nom du rapport']].value,
            'report_date': '%s-%s-%s' % (d[6:10], d[3:5], d[0:2]),
            'file_name': os.path.basename(path),
            'file_data': base64.b64encode(open(path, 'rb').read()).decode()})
        r[H['Statut upload']].value = 'UPLOADE' if st in (200, 201) else 'ERREUR %s: %s' % (st, resp.get('error') or resp)
        print(isin, r[H['Statut upload']].value, flush=True)
        n += 1
        wb.save(MANIFEST)
    wb.save(MANIFEST)


if __name__ == '__main__':
    a = sys.argv[1:]
    if not a:
        sys.exit(__doc__)
    if a[0] == 'login': cmd_login()
    elif a[0] == 'verify': cmd_verify(a[1])
    elif a[0] == 'plan': cmd_plan()
    elif a[0] == 'upload':
        lim = int(a[a.index('--limit') + 1]) if '--limit' in a else None
        cmd_upload('--include-flagged' in a, lim)
    else:
        sys.exit(__doc__)
